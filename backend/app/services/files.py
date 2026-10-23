"""Извлечение текста из загруженных файлов и выгрузка в .docx."""

import io
import re

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader

MAX_BYTES = 10 * 1024 * 1024
# XML не допускает управляющие символы, а в тексте из PDF они встречаются.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    if name.endswith(".pdf"):
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages).strip()
    if name.endswith((".txt", ".md", ".csv", ".tsv")):
        for enc in ("utf-8-sig", "cp1251"):
            try:
                return data.decode(enc).strip()
            except UnicodeDecodeError:
                continue
    raise ValueError("Поддерживаются файлы .docx, .pdf, .txt, .csv")


def to_docx(text: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    for line in text.splitlines():
        doc.add_paragraph(_CONTROL.sub("", line))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
